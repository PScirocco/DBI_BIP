"""理論確認テスト（3色円）の理論値を Excel の数式で計算し、bisim の出力と比べるブックを作る。

Python の理論計算（tools/rgbcircle_theory.py）とは別に、劣化モデルの式を Excel の数式だけで
1フレームずつ計算する。比較には、30フレームすべてで画素値が理想どおり（その色＝255、ほかの色＝0）
の画素を使う（MP4 の非可逆圧縮の影響を受けない画素なので、Excel 側の入力を 255 の定数にできる）。

シート:
  説明       理論式（文字）、記号の意味、比較に使う画素、判定基準、使い方
  パラメータ  モデルパラメータ・Sim条件（青字が入力。書き換えると全体が再計算される）
  計算       1行=1フレーム（0〜1500＝50時間）。見出しの下に各列の Excel 数式を文字で表示（FORMULATEXT）
  比較       7つの評価時点 × R/G/B で、Excel の値と bisim の値（CSV・評価画像）を並べて判定

使い方:
  python tools/rgbcircle_excel.py [--bisim-out <bisimの出力フォルダ>] [--out <xlsx>]
  作成後、Excel で開くと計算される（このスクリプトは Excel があれば開いて計算・保存まで行う）
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
FONT = "Meiryo UI"

SEQ = "rgbcircle"
EVALS = [("02", "eval-0h30m", 0.5), ("04", "eval-1h", 1), ("06", "eval-2h", 2), ("08", "eval-5h", 5),
         ("10", "eval-10h", 10), ("12", "eval-20h", 20), ("14", "eval-50h", 50)]
# 30フレームすべてで画素値が理想どおりの画素（円の中心に近いもの）。(x, y)
PIXELS = {"R": (120, 159), "G": (64, 285), "B": (175, 287)}
COLORS = ["R", "G", "B"]
BGR = {"R": 2, "G": 1, "B": 0}
MODEL = {"N": [1.5, 1.3, 1.7], "K0": [500.0, 1000.0, 100.0], "Q": [1000.0, 800.0, 600.0],
         "B0": [0.7, 0.8, 0.4], "A": [0.0, 0.0, 0.0]}
FRAMES = 1500                                   # 50h ÷ 120秒/フレーム

thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


INPUT = f(color="0000FF")


def cell(ws, ref, value, font=None, fill=None, fmt=None, align=None, border=True):
    c = ws[ref]
    c.value = value
    c.font = font or f()
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


# --------------------------------------------------------------------------- #
def sheet_params(wb):
    ws = wb.create_sheet("パラメータ")
    ws.column_dimensions["A"].width = 26
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 62
    cell(ws, "A1", "パラメータ（青字が入力。書き換えると「計算」「比較」シートが再計算される）",
         f(bold=True, size=11), border=False)
    for col, h in zip("ABCDEF", ["項目", "R", "G", "B", "単位", "説明"]):
        cell(ws, f"{col}3", h, f(bold=True), HEAD_FILL, align=Alignment(horizontal="center"))
    notes = {
        "N": ("−", "加速係数（輝度に対するストレスの効き方）。lum^N"),
        "K0": ("−", "τ（寿命の目安）に相当する係数。大きいほどストレスが貯まりにくい"),
        "Q": ("K", "温度依存の係数（アレニウス型）。exp(Q÷T)"),
        "B0": ("−", "劣化カーブの形 β の基準値"),
        "A": ("1/K", "β の温度依存（β = B0 + A×T）。既定 0"),
    }
    r = 4
    for key in ["N", "K0", "Q", "B0", "A"]:
        cell(ws, f"A{r}", key, f(bold=True))
        for i, col in enumerate("BCD"):
            cell(ws, f"{col}{r}", MODEL[key][i], INPUT)
        cell(ws, f"E{r}", notes[key][0], align=Alignment(horizontal="center"))
        cell(ws, f"F{r}", notes[key][1])
        r += 1
    # r = 9
    cell(ws, "A9", "β（= B0 + A×T）", f(bold=True))
    for col in "BCD":
        cell(ws, f"{col}9", f"={col}7+{col}8*$B$15", fmt="0.000")
    cell(ws, "E9", "−", align=Alignment(horizontal="center"))
    cell(ws, "F9", "劣化カーブの形（A=0 なら B0 と同じ）")

    cell(ws, "A11", "共通（Sim条件）", f(bold=True, size=11), border=False)
    rows = [
        ("ACCEL_RATIO", 3600, "倍", "シミュレーションの加速（1秒の表示を何秒分の Aging とみなすか）", True, "0"),
        ("fps", 30, "1/s", "動画のフレームレート（rgbcircle-aging_movie.mp4）", True, "0"),
        ("固定温度", 25, "℃", "ヒートマップなし（パネル全面が同じ温度）", True, "0.0"),
        ("T（絶対温度）", "=B14+273", "K", "T = 固定温度 + 273（bisim・source/main.py と同じく 273 を足す）", False, "0.0"),
        ("Δt（1フレームの時間）", "=1/B13", "s", "Δt = 1 ÷ fps", False, "0.0000000"),
        ("1フレームの加速込み時間", "=B16*B12", "s", "Δt × ACCEL_RATIO（= 120 秒）", False, "0.0"),
        ("画素値 v", 255, "−", "比較に使う画素の値（その色の円の中＝255）", True, "0"),
        ("ガンマ γ", 2.2, "−", "画素値→輝度の変換 lum = (v/255)^γ", True, "0.0"),
    ]
    r = 12
    for name, val, unit, note, is_input, fmt in rows:
        cell(ws, f"A{r}", name, f(bold=True))
        cell(ws, f"B{r}", val, INPUT if is_input else f(), fmt=fmt)
        cell(ws, f"E{r}", unit, align=Alignment(horizontal="center"))
        cell(ws, f"F{r}", note)
        r += 1
    return ws


P = "パラメータ"
# パラメータシートの参照（列 B/C/D = R/G/B）
PCOL = {"R": "B", "G": "C", "B": "D"}


def pref(key, color):
    row = {"N": 4, "K0": 5, "Q": 6, "B0": 7, "A": 8, "beta": 9}[key]
    return f"{P}!${PCOL[color]}${row}"


ACCEL, FPS, TK, DT, DTACC, V, GAMMA = (f"{P}!$B${r}" for r in (12, 13, 15, 16, 17, 18, 19))


# --------------------------------------------------------------------------- #
def sheet_calc(wb):
    ws = wb.create_sheet("計算")
    first = 5                                   # フレーム0 の行
    ws.freeze_panes = "D5"
    cols = {"k": "A", "sec": "B", "hour": "C"}
    for i, c in enumerate(COLORS):
        base = 4 + i * 4                        # D,H,L
        cols[c] = [get_column_letter(base + j) for j in range(4)]   # deg開始, lum, stat, deg終了
    widths = {"A": 8, "B": 14, "C": 10}
    for c in COLORS:
        for j, col in enumerate(cols[c]):
            widths[col] = 24
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # 1行目: 色のグループ見出し / 2行目: 列名 / 3行目: 数式（文字） / 4行目: 数式の意味
    cell(ws, "A1", "共通", f(bold=True), HEAD_FILL, align=Alignment(horizontal="center"))
    ws.merge_cells("A1:C1")
    names = ["輝度残存率 deg（フレーム開始時）", "輝度 lum", "累積時間 stat（フレーム終了時）", "輝度残存率 deg（フレーム終了時）"]
    meaning = ["exp(−stat(k−1)^β)", "(v/255)^γ × deg開始", "stat(k−1) + Δt×ACCEL×lum^N ÷ (K0×exp(Q÷T))",
               "exp(−stat(k)^β)"]
    for c in COLORS:
        a, b = cols[c][0], cols[c][-1]
        cell(ws, f"{a}1", f"{c}（画素 x={PIXELS[c][0]}, y={PIXELS[c][1]}）", f(bold=True), HEAD_FILL,
             align=Alignment(horizontal="center"))
        ws.merge_cells(f"{a}1:{b}1")
    for col, h in [("A", "フレーム k"), ("B", "累積Aging時間 [s]"), ("C", "[h]")]:
        cell(ws, f"{col}2", h, f(bold=True), HEAD_FILL, align=Alignment(horizontal="center", wrap_text=True))
    for c in COLORS:
        for j, col in enumerate(cols[c]):
            cell(ws, f"{col}2", names[j], f(bold=True), HEAD_FILL,
                 align=Alignment(horizontal="center", wrap_text=True))
    # 3行目：実際の Excel 数式（フレーム1 の行）を文字で表示
    for col in ["B", "C"] + [x for c in COLORS for x in cols[c]]:
        cell(ws, f"{col}3", f"=_xlfn.FORMULATEXT({col}{first + 1})", f(color="7F7F7F", size=8, italic=True),
             SUB_FILL, align=Alignment(wrap_text=True, vertical="top"))
    cell(ws, "A3", "数式", f(bold=True, size=8), SUB_FILL, align=Alignment(vertical="top"))
    cell(ws, "A4", "意味", f(bold=True, size=8), SUB_FILL, align=Alignment(vertical="top"))
    cell(ws, "B4", "k × Δt × ACCEL", f(size=8), SUB_FILL)
    cell(ws, "C4", "[s] ÷ 3600", f(size=8), SUB_FILL)
    for c in COLORS:
        for j, col in enumerate(cols[c]):
            cell(ws, f"{col}4", meaning[j], f(size=8), SUB_FILL, align=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 66
    ws.row_dimensions[4].height = 44

    num = "0.000000000000000"
    # フレーム0（初期状態：ストレス0・劣化なし）
    r0 = first
    cell(ws, f"A{r0}", 0)
    cell(ws, f"B{r0}", 0, fmt="#,##0")
    cell(ws, f"C{r0}", 0, fmt="0.00")
    for c in COLORS:
        d0, lum, st, d1 = cols[c]
        cell(ws, f"{d0}{r0}", "")
        cell(ws, f"{lum}{r0}", "")
        cell(ws, f"{st}{r0}", 0, INPUT, fmt=num)
        cell(ws, f"{d1}{r0}", f"=EXP(-1*{st}{r0}^{pref('beta', c)})", fmt=num)
    for k in range(1, FRAMES + 1):
        r = first + k
        ws[f"A{r}"] = k
        ws[f"B{r}"] = f"=A{r}*{DTACC}"
        ws[f"B{r}"].number_format = "#,##0"
        ws[f"C{r}"] = f"=B{r}/3600"
        ws[f"C{r}"].number_format = "0.00"
        for c in COLORS:
            d0, lum, st, d1 = cols[c]
            ws[f"{d0}{r}"] = f"=EXP(-1*{st}{r - 1}^{pref('beta', c)})"
            ws[f"{lum}{r}"] = f"=({V}/255)^{GAMMA}*{d0}{r}"
            ws[f"{st}{r}"] = (f"={st}{r - 1}+{DT}*{ACCEL}*({lum}{r}^{pref('N', c)})"
                              f"/({pref('K0', c)}*EXP({pref('Q', c)}/{TK}))")
            ws[f"{d1}{r}"] = f"=EXP(-1*{st}{r}^{pref('beta', c)})"
            for col in (d0, lum, st, d1):
                ws[f"{col}{r}"].number_format = num
        for col in ["A", "B", "C"] + [x for c in COLORS for x in cols[c]]:
            ws[f"{col}{r}"].font = f()
    return ws, cols, first


# --------------------------------------------------------------------------- #
def load_bisim(bisim_out: Path) -> dict:
    """比較画素の bisim 値 {(nn, color): (stat, deg, 評価画素値)}。"""
    vals = {}
    for nn, name, _h in EVALS:
        img = cv2.imdecode(np.fromfile(str(bisim_out / f"{SEQ}_{nn}_{name}_image.png"), np.uint8),
                           cv2.IMREAD_COLOR)
        for c in COLORS:
            x, y = PIXELS[c]
            st = np.loadtxt(str(bisim_out / f"{SEQ}_{nn}_{name}_stat_{c.lower()}.csv"), delimiter=",")
            dg = np.loadtxt(str(bisim_out / f"{SEQ}_{nn}_{name}_deg_{c.lower()}.csv"), delimiter=",")
            vals[(nn, c)] = (float(st[y, x]), float(dg[y, x]), int(img[y, x, BGR[c]]))
    return vals


def sheet_compare(wb, cols, first, bisim_out: Path, bisim_label: str):
    ws = wb.create_sheet("比較")
    vals = load_bisim(bisim_out)
    heads = ["時点 [h]", "フレーム数", "色", "画素 (x, y)",
             "累積時間 stat\nExcel", "累積時間 stat\nbisim（CSV）", "相対差",
             "輝度残存率 deg\nExcel", "輝度残存率 deg\nbisim（CSV）", "差",
             "評価画像の画素値\nExcel", "評価画像の画素値\nbisim（PNG）", "差",
             "判定", "bisim の値の出典（CSV は 行=y+1・列=x+1）"]
    widths = [8, 9, 5, 12, 25, 25, 11, 25, 25, 11, 15, 15, 6, 7, 70]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = w
        cell(ws, f"{col}4", h, f(bold=True), HEAD_FILL,
             align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    ws.row_dimensions[4].height = 48
    cell(ws, "A1", "Excel の数式で計算した理論値と bisim の出力の比較", f(bold=True, size=12), border=False)
    cell(ws, "A2", f"bisim の出力：{bisim_label}", f(size=9), border=False)
    num = "0.000000000000000"
    r = 5
    for nn, name, hours in EVALS:
        frames = int(round(hours * 3600 / 120))
        for c in COLORS:
            st_col, d1_col = cols[c][2], cols[c][3]
            x, y = PIXELS[c]
            bst, bdg, bpx = vals[(nn, c)]
            cell(ws, f"A{r}", hours, fmt="0.0")
            cell(ws, f"B{r}", frames)
            cell(ws, f"C{r}", c, align=Alignment(horizontal="center"))
            cell(ws, f"D{r}", f"({x}, {y})", align=Alignment(horizontal="center"))
            cell(ws, f"E{r}", f"=INDEX(計算!{st_col}:{st_col},B{r}+{first})", fmt=num)
            cell(ws, f"F{r}", bst, INPUT, fmt=num)
            cell(ws, f"G{r}", f"=ABS(E{r}-F{r})/E{r}", fmt="0.0E+00")
            cell(ws, f"H{r}", f"=INDEX(計算!{d1_col}:{d1_col},B{r}+{first})", fmt=num)
            cell(ws, f"I{r}", bdg, INPUT, fmt=num)
            cell(ws, f"J{r}", f"=ABS(H{r}-I{r})", fmt="0.0E+00")
            cell(ws, f"K{r}", f"=INT(255*H{r}^(1/{GAMMA}))")
            cell(ws, f"L{r}", bpx, INPUT)
            cell(ws, f"M{r}", f"=ABS(K{r}-L{r})")
            cell(ws, f"N{r}", f'=IF(AND(G{r}<1E-9,J{r}<1E-9,M{r}<=1),"OK","NG")',
                 f(bold=True), align=Alignment(horizontal="center"))
            cell(ws, f"O{r}", f"{SEQ}_{nn}_{name}_stat_{c.lower()}.csv / _deg_{c.lower()}.csv の {y + 1}行{x + 1}列、"
                              f"_image.png の (x={x}, y={y}) の {c}", f(size=8))
            r += 1
    last = r - 1
    cell(ws, "A3", "総合判定", f(bold=True), OK_FILL)
    cell(ws, "B3", f'=IF(COUNTIF(N5:N{last},"NG")=0,"OK（理論どおり）","NG あり")', f(bold=True), OK_FILL)
    ws.merge_cells("B3:D3")
    cell(ws, f"A{last + 2}",
         "判定基準：累積時間の相対差と輝度残存率の差がともに 1E-9 未満、評価画像の画素値の差が 1 以内。"
         "Excel と Python（bisim）は同じ倍精度で計算するが、EXP やべき乗の最後の桁が違うことがあるため、完全一致ではなくこの基準にしている。",
         f(size=9), border=False)
    ws.freeze_panes = "E5"
    return ws


# --------------------------------------------------------------------------- #
def sheet_readme(wb, bisim_label):
    ws = wb.active
    ws.title = "説明"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 120
    lines = [
        ("理論確認テスト（3色円）：Excel の数式による理論値の計算と bisim との比較", "title"),
        ("", None),
        ("■ 目的", "h"),
        ("bisim（劣化シミュレータ）が、劣化モデルの式どおりに累積時間と輝度残存率を計算していることを、"
         "Excel の数式だけで計算した理論値と比べて確かめる。", None),
        ("", None),
        ("■ 劣化モデルの式（暫定モデル。source/main.py の temp_update_stat_and_burn_img と同じ）", "h"),
        ("1フレームごとに、次の (1)〜(3) をこの順に計算する。k はフレーム番号（1, 2, 3, …）、stat(0) = 0。", None),
        ("  (1) フレーム開始時の輝度残存率   deg(k) = exp( −stat(k−1)^β )", "eq"),
        ("  (2) 劣化した後の輝度             lum(k) = (v ÷ 255)^γ × deg(k)", "eq"),
        ("  (3) 累積時間（累積ストレス）     stat(k) = stat(k−1) + Δt × ACCEL_RATIO × lum(k)^N ÷ ( K0 × exp(Q ÷ T) )", "eq"),
        ("  評価（白で表示）したときの画素値   = INT( 255 × deg^(1/γ) )　　※ deg は Aging 終了時点の exp( −stat^β )", "eq"),
        ("  β = B0 + A × T、 T = 温度[℃] + 273、 Δt = 1 ÷ fps", "eq"),
        ("ポイント：(3) では「劣化した後の輝度」でストレスを貯める。劣化が進むほど貯まり方が遅くなるので、"
         "時間×一定量の掛け算にはならず、1フレームずつ順に計算する必要がある（「計算」シートが1行1フレームなのはこのため）。", None),
        ("", None),
        ("■ 記号の意味", "h"),
        ("  stat：累積時間（累積ストレス）  deg：輝度残存率（1＝劣化なし、0 に近いほど劣化）  v：画素値（0〜255）  γ：ガンマ 2.2", None),
        ("  N：加速係数  K0：寿命の目安に相当する係数  Q：温度依存の係数  B0・A：劣化カーブの形 β を決める係数", None),
        ("  ACCEL_RATIO：加速（1秒の表示を何秒分の Aging とみなすか）  fps：動画のフレームレート", None),
        ("", None),
        ("■ 数値例（R、フレーム1）", "h"),
        ("  K0×exp(Q÷T) = 500 × exp(1000 ÷ 298) = 500 × 28.666 = 14332.9", None),
        ("  stat(1) = 0 + (1/30) × 3600 × 1^1.5 ÷ 14332.9 = 120 ÷ 14332.9 = 0.008372", None),
        ("  これを 30 フレーム（1時間）繰り返すと stat = 0.1906、deg = exp(−0.1906^0.7) = 0.7310、評価画素値 = INT(255 × 0.7310^(1/2.2)) = 221", None),
        ("", None),
        ("■ 比較に使う画素", "h"),
        ("  R の円 (x=120, y=159)、G の円 (x=64, y=285)、B の円 (x=175, y=287)。", None),
        ("  Aging 用の動画（rgbcircle-aging_movie.mp4）は MP4 の非可逆圧縮で、画素値が理想の 0/255 から数階調ずれる所がある。"
         "上の画素は 30 フレームすべてで「その色＝255、ほかの色＝0」のまま保たれている画素なので、Excel 側の入力を v = 255 の定数にできる。", None),
        ("  その色以外のチャンネル（例：R の円の中の G・B）は 0 なので劣化しない（このブックでは計算していない）。", None),
        ("", None),
        ("■ シートの内容", "h"),
        ("  パラメータ：モデルパラメータと Sim条件。青字が入力で、書き換えると全体が再計算される", None),
        ("  計算：1行＝1フレーム（0〜1500＝50時間）。3行目に各列の Excel の数式を文字で表示（FORMULATEXT 関数）、4行目に式の意味", None),
        ("  比較：7つの評価時点（0.5h〜50h）× R/G/B で、Excel の値と bisim の値を並べて判定。最上部に総合判定", None),
        ("", None),
        ("■ 比較した bisim の結果", "h"),
        (f"  {bisim_label}", None),
        ("  シーケンス eval_test/rgbcircle/rgbcircle.seq.json（ACCEL_RATIO=3600、固定温度 25℃、Aging→PQ評価 を累積 0.5h〜50h で7回）", None),
        ("", None),
        ("■ 作り直し方", "h"),
        ("  python tools/rgbcircle_excel.py --bisim-out <bisim の出力フォルダ> --out <xlsx>", None),
    ]
    r = 1
    for text, kind in lines:
        c = ws[f"B{r}"]
        c.value = text
        if kind == "title":
            c.font = f(bold=True, size=14)
        elif kind == "h":
            c.font = f(bold=True, size=11)
        elif kind == "eq":
            c.font = Font(name="Consolas", size=11, color="1F3864")
        else:
            c.font = f(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return ws


# --------------------------------------------------------------------------- #
def recalc_with_excel(path: Path) -> bool:
    """Excel があれば開いて全再計算・保存（数式の計算結果をファイルに残す）。"""
    import subprocess
    ps = (f'$xl = New-Object -ComObject Excel.Application; $xl.Visible=$false; $xl.DisplayAlerts=$false; '
          f'$wb = $xl.Workbooks.Open("{path}"); $xl.CalculateFull(); $wb.Worksheets.Item(1).Activate(); '
          f'$wb.Save(); $wb.Close($false); $xl.Quit()')
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True,
                           timeout=300)
        return r.returncode == 0
    except (OSError, subprocess.TimeoutExpired):
        return False


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bisim-out", type=Path, default=REPO / "eval_test_other-PC" / "rgbcircle" / "output")
    ap.add_argument("--out", type=Path, default=REPO / "eval_test" / "rgbcircle" / "理論値_Excel検算.xlsx")
    ap.add_argument("--label", default="別PC（EXE）で実行した結果：eval_test_other-PC/rgbcircle/output")
    a = ap.parse_args()

    wb = Workbook()
    sheet_readme(wb, a.label)
    sheet_params(wb)
    _ws, cols, first = sheet_calc(wb)
    sheet_compare(wb, cols, first, a.bisim_out, a.label)
    out = a.out.resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    ok = recalc_with_excel(out)
    print(f"保存: {out}（Excel で再計算: {'済' if ok else '未（Excel で開くと計算されます）'}）")


if __name__ == "__main__":
    sys.exit(main())
